"""PDF 和 Word 报告生成器。

生成器只消费已经冻结的 JSON 快照，不直接读取数据库，保证文件内容与导出任务
创建时的报告口径一致。依赖在 Worker 环境中加载，API 进程不会因为导出库缺失而
无法提供登录、报告查询等基础接口。
"""

from __future__ import annotations

import html
import os
from datetime import datetime
from pathlib import Path
from typing import Any


SUBJECT_LABELS = {"chinese": "语文", "math": "数学", "english": "英语"}


def _text(value: Any) -> str:
    if value is None or value == "":
        return "—"
    return str(value)


def _mode_label(mode: str) -> str:
    return "教师已复核正式成绩" if mode == "reviewed" else "AI 初评参考（非正式成绩）"


def _font_path() -> Path | None:
    configured = os.getenv("REPORT_FONT_PATH")
    candidates = [
        Path(configured) if configured else None,
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    return next((item for item in candidates if item and item.is_file()), None)


def _snapshot_title(snapshot: dict[str, Any]) -> str:
    batch = snapshot["batch"]
    return f"{batch['title']} 批改报告"


def _display_time(value: Any) -> str:
    raw = _text(value)
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return raw


def _question_comment(question: dict[str, Any]) -> str:
    lines: list[str] = []
    if question.get("teacher_comment"):
        lines.append(f"教师评语：{question['teacher_comment']}")
    if question.get("ai_comment"):
        lines.append(f"AI评语：{question['ai_comment']}")
    deduction = question.get("deduction")
    if isinstance(deduction, list):
        reasons = [str(item.get("reason")) for item in deduction if isinstance(item, dict) and item.get("reason")]
        if reasons:
            lines.append(f"扣分原因：{'；'.join(reasons)}")
    if not lines and question.get("evidence"):
        lines.append(f"证据：{question['evidence']}")
    return "\n".join(lines) or "—"


def build_pdf(snapshot: dict[str, Any], output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_path = _font_path()
    body_font = "Helvetica"
    if font_path:
        pdfmetrics.registerFont(TTFont("PixiaoZhiCJK", str(font_path)))
        body_font = "PixiaoZhiCJK"
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName=body_font, fontSize=18, leading=24, alignment=TA_CENTER, textColor=colors.HexColor("#0f172a"), spaceAfter=10)
    heading_style = ParagraphStyle("ReportHeading", parent=styles["Heading2"], fontName=body_font, fontSize=12, leading=16, textColor=colors.HexColor("#1d4ed8"), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName=body_font, fontSize=9, leading=14, alignment=TA_LEFT, textColor=colors.HexColor("#334155"))
    small_style = ParagraphStyle("ReportSmall", parent=body_style, fontSize=8, leading=11)

    def paragraph(value: Any, style=body_style) -> Paragraph:
        return Paragraph(html.escape(_text(value)).replace("\n", "<br/>") , style)

    def footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont(body_font, 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(18 * mm, 10 * mm, "批小智 · 作业批改报告")
        canvas.drawRightString(192 * mm, 10 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    batch = snapshot["batch"]
    statistics = snapshot["statistics"]
    mode = batch["mode"]
    story: list[Any] = [paragraph(_snapshot_title(snapshot), title_style)]
    if mode == "ai_preview":
        story.append(paragraph("AI 初评参考，非正式成绩。正式成绩必须经过教师逐题复核。", ParagraphStyle("Warning", parent=body_style, textColor=colors.HexColor("#b45309"), backColor=colors.HexColor("#fff7ed"), borderColor=colors.HexColor("#fed7aa"), borderWidth=0.5, borderPadding=6)))
    metadata = [
        [paragraph("批次", small_style), paragraph(batch["title"], small_style), paragraph("班级", small_style), paragraph(batch.get("class_name") or batch.get("class_id"), small_style)],
        [paragraph("学科", small_style), paragraph(SUBJECT_LABELS.get(batch["subject"], batch["subject"]), small_style), paragraph("报告口径", small_style), paragraph(_mode_label(mode), small_style)],
        [paragraph("满分", small_style), paragraph(batch["total_score"], small_style), paragraph("生成时间", small_style), paragraph(_display_time(batch["generated_at"]), small_style)],
    ]
    story.extend([Table(metadata, colWidths=[20 * mm, 65 * mm, 24 * mm, 70 * mm], style=TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])), Spacer(1, 8)])

    story.append(paragraph("班级统计", heading_style))
    stat_rows = [[paragraph("参与学生", small_style), paragraph(f"{statistics['student_count']} 份", small_style), paragraph("平均分", small_style), paragraph(statistics.get("average_score"), small_style)],
                 [paragraph("最高分", small_style), paragraph(statistics.get("highest_score"), small_style), paragraph("待复核", small_style), paragraph(f"{statistics['pending_review_count']} 份", small_style)]]
    story.append(Table(stat_rows, colWidths=[28 * mm, 45 * mm, 28 * mm, 45 * mm], style=TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eff6ff")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#eff6ff")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#bfdbfe")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dbeafe")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])))

    story.append(paragraph("个人成绩单", heading_style))
    score_header = [paragraph(item, small_style) for item in ["学生", "客观题", "主观题", "总分", "状态"]]
    score_rows = [score_header]
    for student in snapshot["students"]:
        score_rows.append([paragraph(student.get("student_name") or student.get("student_code") or "未归属学生", small_style), paragraph(f"{student.get('objective_score') or '—'} / {student.get('objective_max_score') or '—'}", small_style), paragraph(f"{student.get('subjective_score') or '—'} / {student.get('subjective_max_score') or '—'}", small_style), paragraph(f"{student.get('total_score') or '—'} / {student.get('total_max_score') or '—'}", small_style), paragraph("教师已复核" if student.get("status") == "reviewed" else "部分待复核", small_style)])
    score_table = Table(score_rows, colWidths=[34 * mm, 31 * mm, 31 * mm, 31 * mm, 30 * mm], repeatRows=1)
    score_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(score_table)

    if snapshot["students"]:
        story.append(paragraph("逐题明细", heading_style))
    for student in snapshot["students"]:
        story.append(paragraph(f"{student.get('student_name') or student.get('student_code') or '未归属学生'} · 总分 {student.get('total_score') or '—'} / {student.get('total_max_score') or '—'}", ParagraphStyle("StudentHeading", parent=body_style, fontSize=10, leading=14, textColor=colors.HexColor("#0f172a"), spaceBefore=7, spaceAfter=4)))
        rows = [[paragraph(item, small_style) for item in ["题号", "题型", "得分", "评语 / 依据", "状态"]]]
        for question in student.get("questions", []):
            rows.append([paragraph(question.get("question_no"), small_style), paragraph("客观题" if question.get("question_type") == "objective" else "主观题", small_style), paragraph(f"{question.get('score') or '—'} / {question.get('max_score') or '—'}", small_style), paragraph(_question_comment(question), small_style), paragraph("已复核" if question.get("status") == "reviewed" else "待复核", small_style)])
        detail_table = Table(rows, colWidths=[17 * mm, 24 * mm, 25 * mm, 80 * mm, 24 * mm], repeatRows=1)
        detail_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0ecff")), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(detail_table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(str(output_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=_snapshot_title(snapshot), author="批小智")
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def build_docx(snapshot: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    footer = section.footer.paragraphs[0]
    footer.text = "批小智 · 作业批改报告"
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9)
    document.core_properties.title = _snapshot_title(snapshot)
    document.core_properties.author = "批小智"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(_snapshot_title(snapshot))
    run.bold = True
    run.font.size = Pt(18)
    if snapshot["batch"]["mode"] == "ai_preview":
        warning = document.add_paragraph("AI 初评参考，非正式成绩。正式成绩必须经过教师逐题复核。")
        warning.runs[0].bold = True

    batch = snapshot["batch"]
    metadata = document.add_table(rows=3, cols=4)
    metadata.style = "Table Grid"
    values = [["批次", batch["title"], "班级", batch.get("class_name") or batch.get("class_id")], ["学科", SUBJECT_LABELS.get(batch["subject"], batch["subject"]), "报告口径", _mode_label(batch["mode"])], ["满分", batch["total_score"], "生成时间", _display_time(batch["generated_at"])]]
    for row, values_row in zip(metadata.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = _text(value)

    document.add_heading("班级统计", level=1)
    statistics = snapshot["statistics"]
    stat_table = document.add_table(rows=2, cols=4)
    stat_table.style = "Table Grid"
    stats = [["参与学生", f"{statistics['student_count']} 份", "平均分", statistics.get("average_score")], ["最高分", statistics.get("highest_score"), "待复核", f"{statistics['pending_review_count']} 份"]]
    for row, values_row in zip(stat_table.rows, stats):
        for cell, value in zip(row.cells, values_row):
            cell.text = _text(value)

    document.add_heading("个人成绩单", level=1)
    score_table = document.add_table(rows=1, cols=5)
    score_table.style = "Table Grid"
    for cell, value in zip(score_table.rows[0].cells, ["学生", "客观题", "主观题", "总分", "状态"]):
        cell.text = value
    for student in snapshot["students"]:
        row = score_table.add_row().cells
        row[0].text = _text(student.get("student_name") or student.get("student_code") or "未归属学生")
        row[1].text = f"{student.get('objective_score') or '—'} / {student.get('objective_max_score') or '—'}"
        row[2].text = f"{student.get('subjective_score') or '—'} / {student.get('subjective_max_score') or '—'}"
        row[3].text = f"{student.get('total_score') or '—'} / {student.get('total_max_score') or '—'}"
        row[4].text = "教师已复核" if student.get("status") == "reviewed" else "部分待复核"

    document.add_heading("逐题明细", level=1)
    for student in snapshot["students"]:
        document.add_heading(f"{student.get('student_name') or student.get('student_code') or '未归属学生'} · 总分 {student.get('total_score') or '—'} / {student.get('total_max_score') or '—'}", level=2)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ["题号", "题型", "得分", "评语 / 依据", "状态"]):
            cell.text = value
        for question in student.get("questions", []):
            row = table.add_row().cells
            row[0].text = _text(question.get("question_no"))
            row[1].text = "客观题" if question.get("question_type") == "objective" else "主观题"
            row[2].text = f"{question.get('score') or '—'} / {question.get('max_score') or '—'}"
            row[3].text = _question_comment(question)
            row[4].text = "已复核" if question.get("status") == "reviewed" else "待复核"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def build_report_file(snapshot: dict[str, Any], output_path: Path, file_format: str) -> None:
    if file_format == "pdf":
        build_pdf(snapshot, output_path)
    elif file_format == "docx":
        build_docx(snapshot, output_path)
    else:
        raise ValueError(f"unsupported report format: {file_format}")
